# -*- coding: utf-8 -*-
import os
import sys
import datetime
import decimal
import firebirdsql
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Set console encoding
sys.stdout.reconfigure(encoding='utf-8')

# Output directories
DESKTOP_DIR = os.path.join(os.path.expanduser('~'), 'Desktop', 'ORGA_EXPORT_DATA')
PROJECT_DIR = os.path.join(os.getcwd(), 'storage', 'orga_export')

for d in [DESKTOP_DIR, PROJECT_DIR]:
    os.makedirs(d, exist_ok=True)

print(f"Export target folder: {DESKTOP_DIR}")

# Connect to Firebird
conn = firebirdsql.connect(
    host='127.0.0.1',
    database=r'D:\ORGA_SOFT\DATA\ORGA.GDB',
    user='SYSDBA',
    password='masterkey',
    charset='WIN1256'
)
cur = conn.cursor()

# -------------------------------------------------------------
# 1. Fetch Categories
# -------------------------------------------------------------
cur.execute("SELECT PROPRTIES_ID, PROPRTIES_NAME FROM PROPRTIES WHERE PROPRTIES_TYPE = 1")
categories_map = {row[0]: row[1] for row in cur.fetchall() if row[1]}

# -------------------------------------------------------------
# 2. Fetch Stock Balances Map
# -------------------------------------------------------------
cur.execute("""
    SELECT PROD_ID, SUM(TOTAL_QTY_ALL), SUM(TOTAL_COST) 
    FROM STOCK_STOCK 
    GROUP BY PROD_ID
""")
stock_map = {row[0]: {'qty': float(row[1] or 0), 'cost': float(row[2] or 0)} for row in cur.fetchall()}

# -------------------------------------------------------------
# 3. Fetch Accounts Map for Names
# -------------------------------------------------------------
cur.execute("SELECT ACCOUNT_ID, ACCOUNT_NAME, ACCOUNTS_MASTER FROM ACCOUNTS")
accounts_info = {row[0]: {'name': row[1], 'master': row[2]} for row in cur.fetchall()}

# -------------------------------------------------------------
# 4. Prepare Datasets
# -------------------------------------------------------------

# --- Products ---
cur.execute("""
    SELECT 
        p.PROD_ID,
        p.PROD_NAME,
        p.GROUP_TYPE,
        p.BARCODE,
        p.BARCODE_U,
        p.PRIC_COST,
        p.PRICE_1,
        p.PRICE_2,
        p.PRICE_3,
        p.CONSUMER,
        p.QTY_MIN_STOCK,
        p.QTY_MAX_STOCK,
        p.BOXING_B,
        p.NOTS_
    FROM PRODUCTS p
    ORDER BY p.PROD_ID ASC
""")
products_rows = []
for r in cur.fetchall():
    p_id = r[0]
    p_name = r[1]
    grp_id = r[2]
    cat_name = categories_map.get(grp_id, '')
    barcode = r[3] or ''
    barcode_u = r[4] or ''
    cost = float(r[5] or 0)
    p_retail = float(r[6] or 0)
    p_wholesale = float(r[7] or 0)
    p_half_wholesale = float(r[8] or 0)
    p_consumer = float(r[9] or 0)
    min_stock = float(r[10] or 0)
    max_stock = float(r[11] or 0)
    box_qty = float(r[12] or 1)
    notes = r[13] or ''
    
    stk_info = stock_map.get(p_id, {'qty': 0.0, 'cost': 0.0})
    current_stock = stk_info['qty']
    stock_value = stk_info['cost'] if stk_info['cost'] > 0 else (current_stock * cost)
    
    products_rows.append([
        p_id,
        p_name,
        cat_name,
        barcode,
        barcode_u,
        cost,
        p_retail,
        p_wholesale,
        p_half_wholesale,
        p_consumer,
        min_stock,
        max_stock,
        box_qty,
        current_stock,
        stock_value,
        notes
    ])

# --- Chart of Accounts (All) ---
cur.execute("""
    SELECT 
        a.ACCOUNT_ID,
        a.ACCOUNT_NAME,
        m.ACCOUNT_NAME,
        a.ACCOUNT_TEL1,
        a.ACCOUNT_TEL2,
        a.ACCOUNT_ADDRESS,
        a.ACCOUNT_EMAIL,
        a.LIMIT_CASH,
        a.M_S_NET_BALANCE,
        a.B_TAX_ID,
        a.SIGL_ID,
        a.NOTS_
    FROM ACCOUNTS a
    LEFT JOIN ACCOUNTS m ON a.ACCOUNTS_MASTER = m.ACCOUNT_ID
    ORDER BY a.ACCOUNT_ID ASC
""")
all_accounts_rows = []
for r in cur.fetchall():
    all_accounts_rows.append([
        r[0],
        r[1],
        r[2] or 'حساب رئيسي',
        r[3] or '',
        r[4] or '',
        r[5] or '',
        r[6] or '',
        float(r[7] or 0),
        float(r[8] or 0),
        r[9] or '',
        r[10] or '',
        r[11] or ''
    ])

# --- Customers (ACCOUNTS_MASTER = 27) ---
cur.execute("""
    SELECT 
        a.ACCOUNT_ID,
        a.ACCOUNT_NAME,
        a.ACCOUNT_TEL1,
        a.ACCOUNT_TEL2,
        a.ACCOUNT_ADDRESS,
        a.ACCOUNT_EMAIL,
        a.LIMIT_CASH,
        a.M_S_NET_BALANCE,
        a.B_TAX_ID,
        a.SIGL_ID,
        a.NOTS_
    FROM ACCOUNTS a
    WHERE a.ACCOUNTS_MASTER = 27
    ORDER BY a.ACCOUNT_ID ASC
""")
customers_rows = []
for r in cur.fetchall():
    customers_rows.append([
        r[0],
        r[1],
        r[2] or '',
        r[3] or '',
        r[4] or '',
        r[5] or '',
        float(r[6] or 0),
        float(r[7] or 0),
        r[8] or '',
        r[9] or '',
        r[10] or ''
    ])

# --- Suppliers (ACCOUNTS_MASTER = 10) ---
cur.execute("""
    SELECT 
        a.ACCOUNT_ID,
        a.ACCOUNT_NAME,
        a.ACCOUNT_TEL1,
        a.ACCOUNT_TEL2,
        a.ACCOUNT_ADDRESS,
        a.ACCOUNT_EMAIL,
        a.M_S_NET_BALANCE,
        a.B_TAX_ID,
        a.SIGL_ID,
        a.NOTS_
    FROM ACCOUNTS a
    WHERE a.ACCOUNTS_MASTER = 10
    ORDER BY a.ACCOUNT_ID ASC
""")
suppliers_rows = []
for r in cur.fetchall():
    suppliers_rows.append([
        r[0],
        r[1],
        r[2] or '',
        r[3] or '',
        r[4] or '',
        r[5] or '',
        float(r[6] or 0),
        r[7] or '',
        r[8] or '',
        r[9] or ''
    ])

# --- Detailed Stock (STOCK_STOCK) ---
cur.execute("""
    SELECT 
        s.STOCK_ID,
        s.PROD_ID,
        p.PROD_NAME,
        'الرئيسي',
        s.TOTAL_QTY_ALL,
        s.COST,
        s.TOTAL_COST,
        s.PRICE_CONSUMER,
        s.EXPIRE_DATE,
        s.DATE_IN,
        s.NOTS_
    FROM STOCK_STOCK s
    LEFT JOIN PRODUCTS p ON s.PROD_ID = p.PROD_ID
    ORDER BY s.STOCK_ID ASC
""")
stock_rows = []
for r in cur.fetchall():
    stock_rows.append([
        r[0],
        r[1],
        r[2] or '',
        r[3] or '',
        float(r[4] or 0),
        float(r[5] or 0),
        float(r[6] or 0),
        float(r[7] or 0),
        str(r[8]) if r[8] else '',
        str(r[9]) if r[9] else '',
        r[10] or ''
    ])

# --- Sales Invoices Header (INVOICES_H) ---
cur.execute("""
    SELECT 
        h.INVOICES_H_ID,
        h.DATE_D,
        h.TIME_T,
        h.ACCOUNT_ID,
        a.ACCOUNT_NAME,
        h.TOTAL_,
        h.TOTAL_DISCOUNT1,
        h.TOTAL_TOTAL,
        h.INCOME_MONY,
        h.USERS_NAME,
        h.NOTS_
    FROM INVOICES_H h
    LEFT JOIN ACCOUNTS a ON h.ACCOUNT_ID = a.ACCOUNT_ID
    ORDER BY h.INVOICES_H_ID ASC
""")
invoices_h_rows = []
for r in cur.fetchall():
    invoices_h_rows.append([
        r[0],
        str(r[1]) if r[1] else '',
        str(r[2]) if r[2] else '',
        r[3],
        r[4] or 'بيع مباشر',
        float(r[5] or 0),
        float(r[6] or 0),
        float(r[7] or 0),
        float(r[8] or 0),
        r[9] or '',
        r[10] or ''
    ])

# --- Sales Invoice Details (INVOICES_D) ---
cur.execute("""
    SELECT 
        d.INVOICES_D_ID,
        d.INVOICES_H_ID,
        d.DATE_D,
        d.PROD_ID,
        p.PROD_NAME,
        d.TOTAL_QTY_ALL,
        d.PRICE_1,
        d.TOTAL_TOTAL,
        d.COST,
        d.TOTAL_COST,
        d.USERS_NAME
    FROM INVOICES_D d
    LEFT JOIN PRODUCTS p ON d.PROD_ID = p.PROD_ID
    ORDER BY d.INVOICES_D_ID ASC
""")
invoices_d_rows = []
for r in cur.fetchall():
    invoices_d_rows.append([
        r[0],
        r[1],
        str(r[2]) if r[2] else '',
        r[3],
        r[4] or '',
        float(r[5] or 0),
        float(r[6] or 0),
        float(r[7] or 0),
        float(r[8] or 0),
        float(r[9] or 0),
        r[10] or ''
    ])

# --- Purchase Invoices Header (ORDER_H) ---
cur.execute("""
    SELECT 
        h.ORDER_H_ID,
        h.DATE_D,
        h.TIME_T,
        h.ACCOUNT_ID,
        a.ACCOUNT_NAME,
        h.TOTAL_,
        h.TOTAL_TOTAL,
        h.USERS_NAME,
        h.NOTS_
    FROM ORDER_H h
    LEFT JOIN ACCOUNTS a ON h.ACCOUNT_ID = a.ACCOUNT_ID
    ORDER BY h.ORDER_H_ID ASC
""")
order_h_rows = []
for r in cur.fetchall():
    order_h_rows.append([
        r[0],
        str(r[1]) if r[1] else '',
        str(r[2]) if r[2] else '',
        r[3],
        r[4] or '',
        float(r[5] or 0),
        float(r[6] or 0),
        r[7] or '',
        r[8] or ''
    ])

# --- Purchase Invoice Details (ORDER_D) ---
cur.execute("""
    SELECT 
        d.ORDER_D_ID,
        d.ORDER_H_ID,
        d.DATE_D,
        d.PROD_ID,
        p.PROD_NAME,
        d.TOTAL_QTY_ALL,
        d.COST,
        d.TOTAL_COST,
        d.PRICE_1,
        d.EXPIRE_DATE
    FROM ORDER_D d
    LEFT JOIN PRODUCTS p ON d.PROD_ID = p.PROD_ID
    ORDER BY d.ORDER_D_ID ASC
""")
order_d_rows = []
for r in cur.fetchall():
    order_d_rows.append([
        r[0],
        r[1],
        str(r[2]) if r[2] else '',
        r[3],
        r[4] or '',
        float(r[5] or 0),
        float(r[6] or 0),
        float(r[7] or 0),
        float(r[8] or 0),
        str(r[9]) if r[9] else ''
    ])

# --- Cash Receipts (INCOME_CASH) ---
cur.execute("""
    SELECT 
        c.INCOME_CASH_ID,
        c.DATE_D,
        c.TIME_T,
        a.ACCOUNT_NAME,
        c.CASH,
        c.COMMENT,
        c.USERS_NAME
    FROM INCOME_CASH c
    LEFT JOIN ACCOUNTS a ON c.ACCOUNT_ID = a.ACCOUNT_ID
    ORDER BY c.INCOME_CASH_ID ASC
""")
income_cash_rows = []
for r in cur.fetchall():
    income_cash_rows.append([
        r[0],
        str(r[1]) if r[1] else '',
        str(r[2]) if r[2] else '',
        r[3] or '',
        float(r[4] or 0),
        r[5] or '',
        r[6] or ''
    ])

# --- Cash Payments (PAY_CASH) ---
cur.execute("""
    SELECT 
        p.PAY_CASH_ID,
        p.DATE_D,
        p.TIME_T,
        a.ACCOUNT_NAME,
        p.IN_CASH,
        p.NOTS_,
        p.USERS_NAME
    FROM PAY_CASH p
    LEFT JOIN ACCOUNTS a ON p.ACCOUNT_ID = a.ACCOUNT_ID
    ORDER BY p.PAY_CASH_ID ASC
""")
pay_cash_rows = []
for r in cur.fetchall():
    pay_cash_rows.append([
        r[0],
        str(r[1]) if r[1] else '',
        str(r[2]) if r[2] else '',
        r[3] or '',
        float(r[4] or 0),
        r[5] or '',
        r[6] or ''
    ])

conn.close()

# -------------------------------------------------------------
# 5. Build Styled Excel Workbook
# -------------------------------------------------------------
wb = openpyxl.Workbook()
wb.remove(wb.active)  # remove default sheet

# Excel styling definitions
header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
data_font = Font(name="Calibri", size=10)
border_thin = Side(border_style="thin", color="D9D9D9")
cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)
zebra_fill = PatternFill(start_color="F2F7FA", end_color="F2F7FA", fill_type="solid")

sheets_definition = [
    {
        'title': 'الأصناف والمنتجات',
        'headers': ['كود الصنف', 'اسم الصنف', 'المجموعة/التصنيف', 'الباركود', 'باركود إضافي', 'سعر التكلفة', 'سعر قطاعي', 'سعر جملة', 'سعر نصف جملة', 'سعر المستهلك', 'حد الطلب', 'الحد الأقصى', 'القطع بالعلبة', 'الرصيد الحالي', 'قيمة المخزون', 'ملاحظات'],
        'data': products_rows
    },
    {
        'title': 'العملاء',
        'headers': ['كود العميل', 'اسم العميل', 'الهاتف 1', 'الهاتف 2', 'العنوان', 'البريد الإلكتروني', 'حد الائتمان', 'الرصيد الحالي', 'الرقم الضريبي', 'السجل التجاري', 'ملاحظات'],
        'data': customers_rows
    },
    {
        'title': 'الموردين',
        'headers': ['كود المورد', 'اسم المورد', 'الهاتف 1', 'الهاتف 2', 'العنوان', 'البريد الإلكتروني', 'الرصيد الحالي', 'الرقم الضريبي', 'السجل التجاري', 'ملاحظات'],
        'data': suppliers_rows
    },
    {
        'title': 'أرصدة المخزون بالتفصيل',
        'headers': ['كود الباتش', 'كود الصنف', 'اسم الصنف', 'المخزن', 'الكمية المتوفرة', 'سعر التكلفة', 'إجمالي التكلفة', 'سعر المستهلك', 'تاريخ الصلاحية', 'تاريخ الإدخال', 'ملاحظات'],
        'data': stock_rows
    },
    {
        'title': 'فواتير المبيعات',
        'headers': ['رقم الفاتورة', 'التاريخ', 'الوقت', 'كود العميل', 'اسم العميل', 'إجمالي الفاتورة', 'الخصم', 'الصافي المطلوب', 'المدفوع نقداً', 'المستخدم', 'ملاحظات'],
        'data': invoices_h_rows
    },
    {
        'title': 'بنود فواتير المبيعات',
        'headers': ['كود البند', 'رقم الفاتورة', 'التاريخ', 'كود الصنف', 'اسم الصنف', 'الكمية المباعة', 'سعر البيع', 'إجمالي البند', 'التكلفة للوحدة', 'إجمالي التكلفة', 'المستخدم'],
        'data': invoices_d_rows
    },
    {
        'title': 'فواتير المشتريات',
        'headers': ['رقم أمر الشراء', 'التاريخ', 'الوقت', 'كود المورد', 'اسم المورد', 'إجمالي الفاتورة', 'الصافي', 'المستخدم', 'ملاحظات'],
        'data': order_h_rows
    },
    {
        'title': 'بنود فواتير المشتريات',
        'headers': ['كود البند', 'رقم أمر الشراء', 'التاريخ', 'كود الصنف', 'اسم الصنف', 'الكمية المشتراة', 'سعر الشراء', 'إجمالي البند', 'سعر البيع المقترح', 'تاريخ الصلاحية'],
        'data': order_d_rows
    },
    {
        'title': 'سندات القبض (الخزينة)',
        'headers': ['رقم السند', 'التاريخ', 'الوقت', 'اسم الحساب / العميل', 'المبلغ المقبوض', 'البيان', 'المستخدم'],
        'data': income_cash_rows
    },
    {
        'title': 'سندات الصرف (الخزينة)',
        'headers': ['رقم السند', 'التاريخ', 'الوقت', 'اسم الحساب / المورد', 'المبلغ المصروف', 'البيان', 'المستخدم'],
        'data': pay_cash_rows
    },
    {
        'title': 'شجرة الحسابات بالكامل',
        'headers': ['كود الحساب', 'اسم الحساب', 'الحساب الرئيسي', 'الهاتف 1', 'الهاتف 2', 'العنوان', 'البريد', 'حد الائتمان', 'الرصيد الحالي', 'الرقم الضريبي', 'السجل التجاري', 'ملاحظات'],
        'data': all_accounts_rows
    }
]

for s_def in sheets_definition:
    ws = wb.create_sheet(title=s_def['title'])
    ws.sheet_view.rightToLeft = True  # RTL for Arabic
    
    # Headers
    ws.append(s_def['headers'])
    header_row = ws[1]
    for cell in header_row:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28
    
    # Data rows
    row_idx = 2
    for row_data in s_def['data']:
        ws.append(row_data)
        current_row = ws[row_idx]
        use_zebra = (row_idx % 2 == 0)
        for cell in current_row:
            cell.font = data_font
            cell.border = cell_border
            if use_zebra:
                cell.fill = zebra_fill
            # Format numbers
            if isinstance(cell.value, float):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="center" if str(cell.value).isdigit() else "right", vertical="center")
        ws.row_dimensions[row_idx].height = 20
        row_idx += 1
        
    # Auto-fit columns
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

# Save Excel files
excel_desktop = os.path.join(DESKTOP_DIR, 'ORGA_FULL_DATABASE.xlsx')
excel_project = os.path.join(PROJECT_DIR, 'ORGA_FULL_DATABASE.xlsx')
wb.save(excel_desktop)
wb.save(excel_project)
print(f"Saved Excel: {excel_desktop}")

# -------------------------------------------------------------
# 6. Save Separate CSV Files (with UTF-8 BOM for Excel compatibility)
# -------------------------------------------------------------
import csv

for s_def in sheets_definition:
    safe_name = s_def['title'].replace('/', '_').replace(' ', '_')
    csv_desktop = os.path.join(DESKTOP_DIR, f"{safe_name}.csv")
    csv_project = os.path.join(PROJECT_DIR, f"{safe_name}.csv")
    for fpath in [csv_desktop, csv_project]:
        with open(fpath, 'w', encoding='utf-8-sig', newline='') as f:
            writer = csv.writer(f)
            writer.writerow(s_def['headers'])
            for r in s_def['data']:
                writer.writerow(r)

print("Saved CSV files successfully.")

# -------------------------------------------------------------
# 7. Generate Full MySQL-Compatible SQL Dump
# -------------------------------------------------------------
sql_desktop = os.path.join(DESKTOP_DIR, 'ORGA_FULL_DUMP.sql')
sql_project = os.path.join(PROJECT_DIR, 'ORGA_FULL_DUMP.sql')

def escape_sql(val):
    if val is None:
        return 'NULL'
    if isinstance(val, (int, float, decimal.Decimal)):
        return str(val)
    # String
    s = str(val).replace('\\', '\\\\').replace("'", "''").replace('\r', '').replace('\n', '\\n')
    return f"'{s}'"

with open(sql_desktop, 'w', encoding='utf-8') as f:
    f.write("-- ============================================================\n")
    f.write("-- ORGA SOFT (ORGA_10_M) DATABASE EXPORT DUMP\n")
    f.write(f"-- Date: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
    f.write("-- Compatible with: MySQL 5.7+ / MySQL 8.0+ / MariaDB\n")
    f.write("-- Character Set: UTF8MB4\n")
    f.write("-- ============================================================\n\n")
    f.write("SET FOREIGN_KEY_CHECKS = 0;\n")
    f.write("SET NAMES utf8mb4;\n\n")
    
    # 1. Categories
    f.write("-- ---------------------------------------------------------\n")
    f.write("-- Table: orga_categories\n")
    f.write("-- ---------------------------------------------------------\n")
    f.write("DROP TABLE IF EXISTS `orga_categories`;\n")
    f.write("CREATE TABLE `orga_categories` (\n")
    f.write("  `id` INT NOT NULL PRIMARY KEY,\n")
    f.write("  `name` VARCHAR(255) NOT NULL\n")
    f.write(") ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;\n\n")
    if categories_map:
        f.write("INSERT INTO `orga_categories` (`id`, `name`) VALUES\n")
        cat_items = list(categories_map.items())
        for idx, (cid, cname) in enumerate(cat_items):
            sep = ',' if idx < len(cat_items) - 1 else ';'
            f.write(f"({cid}, {escape_sql(cname)}){sep}\n")
        f.write("\n")

    # 2. Products
    f.write("-- ---------------------------------------------------------\n")
    f.write("-- Table: orga_products\n")
    f.write("-- ---------------------------------------------------------\n")
    f.write("DROP TABLE IF EXISTS `orga_products`;\n")
    f.write("CREATE TABLE `orga_products` (\n")
    f.write("  `id` INT NOT NULL PRIMARY KEY,\n")
    f.write("  `name` VARCHAR(255) NOT NULL,\n")
    f.write("  `category_name` VARCHAR(255) DEFAULT NULL,\n")
    f.write("  `barcode` VARCHAR(100) DEFAULT NULL,\n")
    f.write("  `barcode_u` VARCHAR(100) DEFAULT NULL,\n")
    f.write("  `cost_price` DECIMAL(15, 3) DEFAULT 0,\n")
    f.write("  `retail_price` DECIMAL(15, 3) DEFAULT 0,\n")
    f.write("  `wholesale_price` DECIMAL(15, 3) DEFAULT 0,\n")
    f.write("  `half_wholesale_price` DECIMAL(15, 3) DEFAULT 0,\n")
    f.write("  `consumer_price` DECIMAL(15, 3) DEFAULT 0,\n")
    f.write("  `min_stock` DECIMAL(15, 3) DEFAULT 0,\n")
    f.write("  `max_stock` DECIMAL(15, 3) DEFAULT 0,\n")
    f.write("  `box_qty` DECIMAL(15, 3) DEFAULT 1,\n")
    f.write("  `current_stock` DECIMAL(15, 3) DEFAULT 0,\n")
    f.write("  `stock_value` DECIMAL(15, 3) DEFAULT 0,\n")
    f.write("  `notes` TEXT DEFAULT NULL\n")
    f.write(") ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;\n\n")
    if products_rows:
        f.write("INSERT INTO `orga_products` (`id`, `name`, `category_name`, `barcode`, `barcode_u`, `cost_price`, `retail_price`, `wholesale_price`, `half_wholesale_price`, `consumer_price`, `min_stock`, `max_stock`, `box_qty`, `current_stock`, `stock_value`, `notes`) VALUES\n")
        for idx, r in enumerate(products_rows):
            sep = ',' if idx < len(products_rows) - 1 else ';'
            vals = [escape_sql(x) for x in r]
            f.write(f"({', '.join(vals)}){sep}\n")
        f.write("\n")

    # 3. Accounts / Customers & Suppliers
    f.write("-- ---------------------------------------------------------\n")
    f.write("-- Table: orga_accounts\n")
    f.write("-- ---------------------------------------------------------\n")
    f.write("DROP TABLE IF EXISTS `orga_accounts`;\n")
    f.write("CREATE TABLE `orga_accounts` (\n")
    f.write("  `id` INT NOT NULL PRIMARY KEY,\n")
    f.write("  `name` VARCHAR(255) NOT NULL,\n")
    f.write("  `parent_account` VARCHAR(255) DEFAULT NULL,\n")
    f.write("  `phone1` VARCHAR(50) DEFAULT NULL,\n")
    f.write("  `phone2` VARCHAR(50) DEFAULT NULL,\n")
    f.write("  `address` VARCHAR(255) DEFAULT NULL,\n")
    f.write("  `email` VARCHAR(100) DEFAULT NULL,\n")
    f.write("  `limit_cash` DECIMAL(15, 3) DEFAULT 0,\n")
    f.write("  `current_balance` DECIMAL(15, 3) DEFAULT 0,\n")
    f.write("  `tax_id` VARCHAR(50) DEFAULT NULL,\n")
    f.write("  `commercial_reg` VARCHAR(50) DEFAULT NULL,\n")
    f.write("  `notes` TEXT DEFAULT NULL\n")
    f.write(") ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;\n\n")
    if all_accounts_rows:
        f.write("INSERT INTO `orga_accounts` (`id`, `name`, `parent_account`, `phone1`, `phone2`, `address`, `email`, `limit_cash`, `current_balance`, `tax_id`, `commercial_reg`, `notes`) VALUES\n")
        for idx, r in enumerate(all_accounts_rows):
            sep = ',' if idx < len(all_accounts_rows) - 1 else ';'
            vals = [escape_sql(x) for x in r]
            f.write(f"({', '.join(vals)}){sep}\n")
        f.write("\n")

    # 4. Detailed Stock
    f.write("-- ---------------------------------------------------------\n")
    f.write("-- Table: orga_stock_details\n")
    f.write("-- ---------------------------------------------------------\n")
    f.write("DROP TABLE IF EXISTS `orga_stock_details`;\n")
    f.write("CREATE TABLE `orga_stock_details` (\n")
    f.write("  `stock_id` INT NOT NULL PRIMARY KEY,\n")
    f.write("  `product_id` INT NOT NULL,\n")
    f.write("  `product_name` VARCHAR(255) DEFAULT NULL,\n")
    f.write("  `store_name` VARCHAR(100) DEFAULT NULL,\n")
    f.write("  `available_qty` DECIMAL(15, 3) DEFAULT 0,\n")
    f.write("  `cost_price` DECIMAL(15, 3) DEFAULT 0,\n")
    f.write("  `total_cost` DECIMAL(15, 3) DEFAULT 0,\n")
    f.write("  `consumer_price` DECIMAL(15, 3) DEFAULT 0,\n")
    f.write("  `expire_date` DATE DEFAULT NULL,\n")
    f.write("  `date_in` DATE DEFAULT NULL,\n")
    f.write("  `notes` TEXT DEFAULT NULL\n")
    f.write(") ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;\n\n")
    if stock_rows:
        f.write("INSERT INTO `orga_stock_details` (`stock_id`, `product_id`, `product_name`, `store_name`, `available_qty`, `cost_price`, `total_cost`, `consumer_price`, `expire_date`, `date_in`, `notes`) VALUES\n")
        for idx, r in enumerate(stock_rows):
            sep = ',' if idx < len(stock_rows) - 1 else ';'
            vals = [escape_sql(x if x != '' else None) for x in r]
            f.write(f"({', '.join(vals)}){sep}\n")
        f.write("\n")

    # 5. Sales Invoices
    f.write("-- ---------------------------------------------------------\n")
    f.write("-- Table: orga_sales_invoices\n")
    f.write("-- ---------------------------------------------------------\n")
    f.write("DROP TABLE IF EXISTS `orga_sales_invoices`;\n")
    f.write("CREATE TABLE `orga_sales_invoices` (\n")
    f.write("  `id` INT NOT NULL PRIMARY KEY,\n")
    f.write("  `invoice_date` DATE DEFAULT NULL,\n")
    f.write("  `invoice_time` VARCHAR(20) DEFAULT NULL,\n")
    f.write("  `customer_id` INT DEFAULT NULL,\n")
    f.write("  `customer_name` VARCHAR(255) DEFAULT NULL,\n")
    f.write("  `total_amount` DECIMAL(15, 3) DEFAULT 0,\n")
    f.write("  `discount_amount` DECIMAL(15, 3) DEFAULT 0,\n")
    f.write("  `net_amount` DECIMAL(15, 3) DEFAULT 0,\n")
    f.write("  `paid_amount` DECIMAL(15, 3) DEFAULT 0,\n")
    f.write("  `created_by` VARCHAR(100) DEFAULT NULL,\n")
    f.write("  `notes` TEXT DEFAULT NULL\n")
    f.write(") ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;\n\n")
    if invoices_h_rows:
        f.write("INSERT INTO `orga_sales_invoices` (`id`, `invoice_date`, `invoice_time`, `customer_id`, `customer_name`, `total_amount`, `discount_amount`, `net_amount`, `paid_amount`, `created_by`, `notes`) VALUES\n")
        for idx, r in enumerate(invoices_h_rows):
            sep = ',' if idx < len(invoices_h_rows) - 1 else ';'
            vals = [escape_sql(x if x != '' else None) for x in r]
            f.write(f"({', '.join(vals)}){sep}\n")
        f.write("\n")

    # 6. Sales Invoice Details
    f.write("-- ---------------------------------------------------------\n")
    f.write("-- Table: orga_sales_invoice_items\n")
    f.write("-- ---------------------------------------------------------\n")
    f.write("DROP TABLE IF EXISTS `orga_sales_invoice_items`;\n")
    f.write("CREATE TABLE `orga_sales_invoice_items` (\n")
    f.write("  `id` INT NOT NULL PRIMARY KEY,\n")
    f.write("  `invoice_id` INT NOT NULL,\n")
    f.write("  `item_date` DATE DEFAULT NULL,\n")
    f.write("  `product_id` INT NOT NULL,\n")
    f.write("  `product_name` VARCHAR(255) DEFAULT NULL,\n")
    f.write("  `quantity` DECIMAL(15, 3) DEFAULT 0,\n")
    f.write("  `unit_price` DECIMAL(15, 3) DEFAULT 0,\n")
    f.write("  `total_price` DECIMAL(15, 3) DEFAULT 0,\n")
    f.write("  `unit_cost` DECIMAL(15, 3) DEFAULT 0,\n")
    f.write("  `total_cost` DECIMAL(15, 3) DEFAULT 0,\n")
    f.write("  `created_by` VARCHAR(100) DEFAULT NULL\n")
    f.write(") ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;\n\n")
    if invoices_d_rows:
        f.write("INSERT INTO `orga_sales_invoice_items` (`id`, `invoice_id`, `item_date`, `product_id`, `product_name`, `quantity`, `unit_price`, `total_price`, `unit_cost`, `total_cost`, `created_by`) VALUES\n")
        for idx, r in enumerate(invoices_d_rows):
            sep = ',' if idx < len(invoices_d_rows) - 1 else ';'
            vals = [escape_sql(x if x != '' else None) for x in r]
            f.write(f"({', '.join(vals)}){sep}\n")
        f.write("\n")

    # 7. Purchase Invoices
    f.write("-- ---------------------------------------------------------\n")
    f.write("-- Table: orga_purchase_invoices\n")
    f.write("-- ---------------------------------------------------------\n")
    f.write("DROP TABLE IF EXISTS `orga_purchase_invoices`;\n")
    f.write("CREATE TABLE `orga_purchase_invoices` (\n")
    f.write("  `id` INT NOT NULL PRIMARY KEY,\n")
    f.write("  `order_date` DATE DEFAULT NULL,\n")
    f.write("  `order_time` VARCHAR(20) DEFAULT NULL,\n")
    f.write("  `supplier_id` INT DEFAULT NULL,\n")
    f.write("  `supplier_name` VARCHAR(255) DEFAULT NULL,\n")
    f.write("  `total_amount` DECIMAL(15, 3) DEFAULT 0,\n")
    f.write("  `net_amount` DECIMAL(15, 3) DEFAULT 0,\n")
    f.write("  `created_by` VARCHAR(100) DEFAULT NULL,\n")
    f.write("  `notes` TEXT DEFAULT NULL\n")
    f.write(") ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;\n\n")
    if order_h_rows:
        f.write("INSERT INTO `orga_purchase_invoices` (`id`, `order_date`, `order_time`, `supplier_id`, `supplier_name`, `total_amount`, `net_amount`, `created_by`, `notes`) VALUES\n")
        for idx, r in enumerate(order_h_rows):
            sep = ',' if idx < len(order_h_rows) - 1 else ';'
            vals = [escape_sql(x if x != '' else None) for x in r]
            f.write(f"({', '.join(vals)}){sep}\n")
        f.write("\n")

    # 8. Purchase Invoice Details
    f.write("-- ---------------------------------------------------------\n")
    f.write("-- Table: orga_purchase_invoice_items\n")
    f.write("-- ---------------------------------------------------------\n")
    f.write("DROP TABLE IF EXISTS `orga_purchase_invoice_items`;\n")
    f.write("CREATE TABLE `orga_purchase_invoice_items` (\n")
    f.write("  `id` INT NOT NULL PRIMARY KEY,\n")
    f.write("  `order_id` INT NOT NULL,\n")
    f.write("  `item_date` DATE DEFAULT NULL,\n")
    f.write("  `product_id` INT NOT NULL,\n")
    f.write("  `product_name` VARCHAR(255) DEFAULT NULL,\n")
    f.write("  `quantity` DECIMAL(15, 3) DEFAULT 0,\n")
    f.write("  `unit_cost` DECIMAL(15, 3) DEFAULT 0,\n")
    f.write("  `total_cost` DECIMAL(15, 3) DEFAULT 0,\n")
    f.write("  `retail_price` DECIMAL(15, 3) DEFAULT 0,\n")
    f.write("  `expire_date` DATE DEFAULT NULL\n")
    f.write(") ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;\n\n")
    if order_d_rows:
        f.write("INSERT INTO `orga_purchase_invoice_items` (`id`, `order_id`, `item_date`, `product_id`, `product_name`, `quantity`, `unit_cost`, `total_cost`, `retail_price`, `expire_date`) VALUES\n")
        for idx, r in enumerate(order_d_rows):
            sep = ',' if idx < len(order_d_rows) - 1 else ';'
            vals = [escape_sql(x if x != '' else None) for x in r]
            f.write(f"({', '.join(vals)}){sep}\n")
        f.write("\n")

    # 9. Cash Income
    f.write("-- ---------------------------------------------------------\n")
    f.write("-- Table: orga_cash_income\n")
    f.write("-- ---------------------------------------------------------\n")
    f.write("DROP TABLE IF EXISTS `orga_cash_income`;\n")
    f.write("CREATE TABLE `orga_cash_income` (\n")
    f.write("  `id` INT NOT NULL PRIMARY KEY,\n")
    f.write("  `receipt_date` DATE DEFAULT NULL,\n")
    f.write("  `receipt_time` VARCHAR(20) DEFAULT NULL,\n")
    f.write("  `account_name` VARCHAR(255) DEFAULT NULL,\n")
    f.write("  `amount` DECIMAL(15, 3) DEFAULT 0,\n")
    f.write("  `notes` TEXT DEFAULT NULL,\n")
    f.write("  `created_by` VARCHAR(100) DEFAULT NULL\n")
    f.write(") ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;\n\n")
    if income_cash_rows:
        f.write("INSERT INTO `orga_cash_income` (`id`, `receipt_date`, `receipt_time`, `account_name`, `amount`, `notes`, `created_by`) VALUES\n")
        for idx, r in enumerate(income_cash_rows):
            sep = ',' if idx < len(income_cash_rows) - 1 else ';'
            vals = [escape_sql(x if x != '' else None) for x in r]
            f.write(f"({', '.join(vals)}){sep}\n")
        f.write("\n")

    # 10. Cash Payments
    f.write("-- ---------------------------------------------------------\n")
    f.write("-- Table: orga_cash_payments\n")
    f.write("-- ---------------------------------------------------------\n")
    f.write("DROP TABLE IF EXISTS `orga_cash_payments`;\n")
    f.write("CREATE TABLE `orga_cash_payments` (\n")
    f.write("  `id` INT NOT NULL PRIMARY KEY,\n")
    f.write("  `payment_date` DATE DEFAULT NULL,\n")
    f.write("  `payment_time` VARCHAR(20) DEFAULT NULL,\n")
    f.write("  `account_name` VARCHAR(255) DEFAULT NULL,\n")
    f.write("  `amount` DECIMAL(15, 3) DEFAULT 0,\n")
    f.write("  `notes` TEXT DEFAULT NULL,\n")
    f.write("  `created_by` VARCHAR(100) DEFAULT NULL\n")
    f.write(") ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;\n\n")
    if pay_cash_rows:
        f.write("INSERT INTO `orga_cash_payments` (`id`, `payment_date`, `payment_time`, `account_name`, `amount`, `notes`, `created_by`) VALUES\n")
        for idx, r in enumerate(pay_cash_rows):
            sep = ',' if idx < len(pay_cash_rows) - 1 else ';'
            vals = [escape_sql(x if x != '' else None) for x in r]
            f.write(f"({', '.join(vals)}){sep}\n")
        f.write("\n")

    f.write("SET FOREIGN_KEY_CHECKS = 1;\n")

# Copy SQL file to project dir
with open(sql_desktop, 'r', encoding='utf-8') as src, open(sql_project, 'w', encoding='utf-8') as dst:
    dst.write(src.read())

print(f"Saved SQL dump: {sql_desktop}")
print("ALL EXPORT OPERATIONS COMPLETED SUCCESSFULLY!")
